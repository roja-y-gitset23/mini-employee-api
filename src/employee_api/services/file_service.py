"""FileProcessor: CSV, Excel and PDF parsing into employee row dictionaries."""

import csv
import io
from typing import Dict, List

import fitz  # PyMuPDF
from openpyxl import load_workbook

from employee_api.utils.decorators import log_execution
from employee_api.utils.logger import get_logger

logger = get_logger(__name__)

EXPECTED_FIELDS = ("name", "email", "department", "position", "salary")


class FileProcessingError(ValueError):
    """Raised when an uploaded file cannot be parsed."""


class FileProcessor:
    @staticmethod
    def _normalise_header(header: object) -> str:
        return str(header).strip().lower().replace(" ", "_") if header is not None else ""

    @staticmethod
    def _require_columns(headers: List[str]) -> None:
        missing = [field for field in EXPECTED_FIELDS if field not in headers]
        if missing:
            raise FileProcessingError(f"Missing required columns: {', '.join(missing)}")

    # ------------------------------------------------------------------ #
    # CSV
    # ------------------------------------------------------------------ #
    @log_execution
    def parse_csv(self, content: bytes) -> List[Dict[str, str]]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise FileProcessingError("CSV file must be UTF-8 encoded") from exc

        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise FileProcessingError("CSV file has no header row")

        headers = [self._normalise_header(h) for h in reader.fieldnames]
        self._require_columns(headers)
        reader.fieldnames = headers

        rows: List[Dict[str, str]] = []
        for raw in reader:
            row = {
                key: (value or "").strip()
                for key, value in raw.items()
                if key in EXPECTED_FIELDS and isinstance(value, (str, type(None)))
            }
            if any(row.values()):
                rows.append(row)
        return rows

    # ------------------------------------------------------------------ #
    # Excel
    # ------------------------------------------------------------------ #
    @log_execution
    def parse_excel(self, content: bytes) -> List[Dict[str, str]]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many different error types
            raise FileProcessingError("Invalid or unreadable Excel file") from exc

        try:
            sheet = workbook.active
            if sheet is None:
                raise FileProcessingError("Excel file contains no worksheet")

            rows_iter = sheet.iter_rows(values_only=True)
            header_row = next(rows_iter, None)
            if not header_row:
                raise FileProcessingError("Excel sheet is empty")

            headers = [self._normalise_header(h) for h in header_row]
            self._require_columns(headers)

            rows: List[Dict[str, str]] = []
            for values in rows_iter:
                row: Dict[str, str] = {}
                for header, value in zip(headers, values):
                    if header in EXPECTED_FIELDS:
                        row[header] = "" if value is None else str(value).strip()
                if any(row.values()):
                    rows.append(row)
            return rows
        finally:
            workbook.close()

    # ------------------------------------------------------------------ #
    # PDF
    # ------------------------------------------------------------------ #
    @log_execution
    def extract_pdf_text(self, content: bytes) -> str:
        try:
            with fitz.open(stream=content, filetype="pdf") as document:
                pages = [page.get_text("text") for page in document]
        except Exception as exc:  # PyMuPDF raises several error types
            raise FileProcessingError("Invalid or unreadable PDF file") from exc

        text = "\n".join(pages).strip()
        if not text:
            raise FileProcessingError("No extractable text found in PDF")
        return text

    @staticmethod
    def parse_employee_lines(text: str) -> List[Dict[str, str]]:
        """Parse lines of: name | email | department | position | salary.

        Pipe, tab or comma delimiters are supported (pipe recommended). Lines that do
        not contain exactly five fields (titles, page numbers) are ignored.
        """
        rows: List[Dict[str, str]] = []
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            delimiter = "|" if "|" in line else ("\t" if "\t" in line else ",")
            parts = [part.strip() for part in line.split(delimiter)]
            if len(parts) != len(EXPECTED_FIELDS):
                continue
            if parts[0].lower() == "name" and parts[1].lower() == "email":
                continue
            row = dict(zip(EXPECTED_FIELDS, parts))
            row["salary"] = row["salary"].replace("$", "").replace(",", "")
            rows.append(row)

        if not rows:
            raise FileProcessingError("No employee records found in PDF text")
        return rows

    @log_execution
    def parse_pdf(self, content: bytes) -> List[Dict[str, str]]:
        return self.parse_employee_lines(self.extract_pdf_text(content))